"""Unit tests for the ingestion pipeline (DOCX, cleaning, chunking, errors)."""
from __future__ import annotations

from pathlib import Path

import pytest

from app.core.ingestion import (
    SUPPORTED_EXTENSIONS,
    CorruptedDocumentError,
    EmptyDocumentError,
    UnsupportedFormatError,
    chunk_text,
    clean_text,
    extract_text,
    ingest_file,
)


def test_clean_text_dehyphenates_and_collapses():
    raw = "This is an exam-\nple   of   noisy\r\n\n\n\ntext."
    cleaned = clean_text(raw)
    assert "example" in cleaned
    assert "  " not in cleaned
    assert "\n\n\n" not in cleaned


def test_chunking_respects_size_and_overlap():
    text = " ".join(f"Sentence number {i}." for i in range(200))
    chunks = chunk_text(text, chunk_size=200, chunk_overlap=40)
    assert len(chunks) > 1
    # Each chunk should be roughly within the target size (allow some slack).
    assert all(len(c.text) <= 260 for c in chunks)
    # Chunk indices are sequential.
    assert [c.chunk_index for c in chunks] == list(range(len(chunks)))


def test_oversized_sentence_is_hard_split():
    text = "x" * 1000
    chunks = chunk_text(text, chunk_size=200, chunk_overlap=20)
    assert len(chunks) > 1
    assert all(len(c.text) <= 200 for c in chunks)


def test_docx_in_supported_extensions():
    assert ".docx" in SUPPORTED_EXTENSIONS


def test_extract_docx_paragraphs_and_tables(tmp_path: Path):
    pytest.importorskip("docx")
    from docx import Document

    path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("PrivateRAG supports Word documents.")
    doc.add_paragraph("Second paragraph about Elastic IPs.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Service"
    table.rows[0].cells[1].text = "EC2"
    doc.save(path)

    pages = extract_text(path)
    assert len(pages) == 1
    assert pages[0][0] is None
    text = pages[0][1]
    assert "PrivateRAG supports Word documents." in text
    assert "Elastic IPs" in text
    assert "Service | EC2" in text

    chunks = ingest_file(path, chunk_size=500, chunk_overlap=50)
    assert len(chunks) >= 1
    assert any("Word documents" in c.text for c in chunks)


def test_empty_docx_raises(tmp_path: Path):
    pytest.importorskip("docx")
    from docx import Document

    path = tmp_path / "blank.docx"
    Document().save(path)
    with pytest.raises(EmptyDocumentError, match="No extractable text"):
        ingest_file(path, chunk_size=500, chunk_overlap=50)


def test_corrupt_docx_raises(tmp_path: Path):
    path = tmp_path / "bad.docx"
    path.write_bytes(b"not-a-real-docx-file")
    with pytest.raises(CorruptedDocumentError):
        extract_text(path)


def test_unsupported_extension_raises(tmp_path: Path):
    path = tmp_path / "notes.xlsx"
    path.write_text("a,b,c\n1,2,3\n", encoding="utf-8")
    with pytest.raises(UnsupportedFormatError, match="Unsupported file type"):
        extract_text(path)


def test_empty_txt_raises(tmp_path: Path):
    path = tmp_path / "empty.txt"
    path.write_text("   \n\n  ", encoding="utf-8")
    with pytest.raises(EmptyDocumentError):
        ingest_file(path, chunk_size=500, chunk_overlap=50)


def test_txt_ingest_ok(tmp_path: Path):
    path = tmp_path / "notes.txt"
    path.write_text("An Elastic IP is a static public IPv4 address.", encoding="utf-8")
    chunks = ingest_file(path, chunk_size=500, chunk_overlap=50)
    assert len(chunks) == 1
    assert "Elastic IP" in chunks[0].text
